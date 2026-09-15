"""Tests for the Phase 2.5 Word authoring workbook."""

import yaml
from docx import Document
from src.label.taxonomy import load_taxonomy
from src.pilot.render_authoring_docx import CASE_GUIDES, render_workbook


def test_render_workbook_structure(tmp_path) -> None:
    taxonomy_path = tmp_path / "taxonomy.yaml"
    intents = []
    for guide in CASE_GUIDES:
        intents.append(
            {
                "id": guide.intent_id,
                "family": "test_family",
                "definition": f"Definition for {guide.intent_id}",
                "positive_examples": ["one", "two"],
                "negative_example": "negative",
                "negative_rationale": "boundary",
            }
        )
    for index in range(4):
        intents.append(
            {
                "id": f"appendix_only_{index}",
                "family": "appendix_family",
                "definition": "Appendix definition",
                "positive_examples": ["one", "two"],
                "negative_example": "negative",
                "negative_rationale": "boundary",
            }
        )
    taxonomy_path.write_text(
        yaml.safe_dump({"version": 1, "intents": intents, "languages": [], "severities": []}),
        encoding="utf-8",
    )
    output = tmp_path / "workbook.docx"

    render_workbook(load_taxonomy(taxonomy_path), output)

    document = Document(output)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert len(CASE_GUIDES) == 20
    assert len(document.tables) == 20
    assert all(
        f"Case {index:02d}: {guide.intent_id}" in text for index, guide in enumerate(CASE_GUIDES, 1)
    )
    assert all(f"appendix_only_{index}" in text for index in range(4))
    assert text.count("Complete example complaints") == 20
    assert all(f"Definition for {guide.intent_id}" in text for guide in CASE_GUIDES)
    assert text.count("one") >= 20
    assert text.count("two") >= 20
    assert "Inspiration scenario (do not copy)" in text
    assert "Do not enter real names, phone numbers, CNICs, or accounts" in text
