"""Render a private Word workbook for human Phase 2.5 complaint authoring."""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from src.label.taxonomy import Intent, Taxonomy, load_taxonomy
from src.pilot.authoring import LANGUAGE_FORMS

DEFAULT_OUTPUT = Path("spike/phase2_5/UrduCX_Phase2_5_Authoring_Workbook.docx")
DEFAULT_TAXONOMY = Path("config/taxonomy.yaml")


@dataclass(frozen=True)
class CaseGuide:
    """One suggested pilot case without writing the four human-owned variants."""

    intent_id: str
    scenario_brief: str
    concrete_fact: str
    boundary_note: str


CASE_GUIDES = (
    CaseGuide(
        "unauthorized_vas_deduction",
        "A weekly game or subscription charge appeared even though the customer says they "
        "never activated it.",
        "Use a specific charge and date, such as Rs 30 and 18 August.",
        "Name the unwanted service; otherwise this may be unexplained balance loss.",
    ),
    CaseGuide(
        "balance_disappeared_unexplained",
        "The customer's balance dropped without any transaction or named service explaining "
        "the loss.",
        "Use a missing amount, such as Rs 200.",
        "Do not name a VAS, transfer, package, or other confirmed cause.",
    ),
    CaseGuide(
        "package_not_activated",
        "Payment for a mobile bundle succeeded, but its minutes or data never became available.",
        "Use a price and package duration, such as Rs 350 and seven days.",
        "The package must fail to activate, not merely expire after working.",
    ),
    CaseGuide(
        "overcharging_dispute",
        "A specific transaction cost more than the tariff shown before confirmation.",
        "Include both advertised and charged amounts, such as Rs 10 and Rs 25.",
        "This is a concrete price discrepancy, not a general complaint that fees are high.",
    ),
    CaseGuide(
        "refund_request",
        "A paid service was not delivered and the customer explicitly asks for their money back.",
        "Use a refund amount, such as Rs 1,500.",
        "The message must explicitly request return of money.",
    ),
    CaseGuide(
        "login_failure",
        "Correct credentials are entered repeatedly, but the app rejects the login attempt.",
        "Use an attempt count or time, such as three attempts since 9 AM.",
        "The blockage is login itself, not a verification code that never arrived.",
    ),
    CaseGuide(
        "otp_not_received",
        "The customer requests a verification code several times, but no OTP arrives.",
        "Use a wait time and attempt count, such as 20 minutes and three requests.",
        "The OTP must be absent, not received and then rejected as invalid.",
    ),
    CaseGuide(
        "kyc_verification_stuck",
        "An identity-verification or account-upgrade request remains pending without an outcome.",
        "Use a pending duration, such as 10 days.",
        "The process remains pending; a completed freeze belongs to account_blocked_frozen.",
    ),
    CaseGuide(
        "account_blocked_frozen",
        "The provider has frozen the account, preventing access to funds or services.",
        "Use a date and trapped balance, such as 15 August and Rs 4,000.",
        "State that the account is blocked or frozen, not merely that login failed.",
    ),
    CaseGuide(
        "transfer_failed_money_deducted",
        "A transfer did not reach the intended recipient even though the sender's balance was "
        "reduced.",
        "Use an amount and transaction reference, such as Rs 5,000 and TXN-2048.",
        "The recipient did not receive it; do not describe a successful transfer to the wrong "
        "person.",
    ),
    CaseGuide(
        "transfer_to_wrong_recipient",
        "The sender entered the wrong number and the transfer reached an unintended recipient.",
        "Use a transfer amount, such as Rs 2,000.",
        "The transfer succeeded technically; it did not simply disappear or fail.",
    ),
    CaseGuide(
        "cash_out_agent_problem",
        "At a physical agent shop, wallet balance was deducted but cash was not handed over.",
        "Use an amount and time, such as Rs 3,000 at 6 PM.",
        "A physical cash-out agent must be involved, not only an in-app transfer.",
    ),
    CaseGuide(
        "bill_payment_not_reflected",
        "A utility bill was paid and deducted in the app, but the biller still shows it as unpaid.",
        "Use a bill amount and date, such as Rs 6,200 on 12 August.",
        "The payment must complete; an app crash before payment is a different failure.",
    ),
    CaseGuide(
        "loan_repayment_dispute",
        "A loan repayment included an unexpected rollover fee or automatic extra deduction.",
        "Use principal and extra-fee amounts, such as Rs 1,000 and Rs 250.",
        "Focus on repayment terms or deductions, not rejection of a new loan application.",
    ),
    CaseGuide(
        "scam_impersonation_report",
        "A caller claims to be support staff and asks the customer to reveal an OTP or PIN.",
        "Use a call time or case reference, such as 4:30 PM or SCAM-17.",
        "The person must be falsely impersonating staff, not a genuine rude support agent.",
    ),
    CaseGuide(
        "fake_payment_screenshot",
        "A buyer shows a forged payment screenshot, receives goods, but no money reaches the "
        "seller.",
        "Use a sale amount and order reference, such as Rs 8,500 and ORD-92.",
        "The screenshot itself is false; this is not a dispute after a genuine payment.",
    ),
    CaseGuide(
        "account_compromised",
        "An unknown person accessed the customer's account and made an unauthorized transaction.",
        "Use a transaction amount and date, such as Rs 7,000 on 20 August.",
        "Describe actual unauthorized access, not a protective block after failed logins.",
    ),
    CaseGuide(
        "network_coverage_complaint",
        "Calls repeatedly drop because there is little or no signal in a named area.",
        "Use a location and duration, such as Gulberg for three days.",
        "This concerns signal availability, not slow data with full signal bars.",
    ),
    CaseGuide(
        "data_speed_complaint",
        "The phone shows a strong 4G connection, but downloads and video remain extremely slow.",
        "Use a measured speed or duration, such as 0.5 Mbps since yesterday.",
        "A connection exists; complete unavailability belongs to service_outage_report.",
    ),
    CaseGuide(
        "service_outage_report",
        "Calls, data, or app services are completely unavailable for multiple people in one area.",
        "Use an outage start time and duration, such as 8 AM for six hours.",
        "Describe a broad outage, not one device or one customer's weak coverage.",
    ),
)

LANGUAGE_LABELS = {
    "urdu_script": "Urdu script",
    "roman_urdu": "Natural Roman Urdu",
    "code_switched": "Urdu-English code-switched",
    "english": "English",
}


def _shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _set_cell_text(cell, text: str, *, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _configure_document(document: DocumentType) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    for style_name, size, color in (
        ("Title", 28, RGBColor(24, 76, 61)),
        ("Heading 1", 20, RGBColor(24, 76, 61)),
        ("Heading 2", 14, RGBColor(42, 90, 74)),
    ):
        style = document.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = color


def _add_intro(document: DocumentType) -> None:
    document.add_heading("UrduCX-Bench Phase 2.5", 0)
    subtitle = document.add_paragraph("Human Authoring Workbook | 20 Script-Gap Pilot Cases")
    subtitle.style = document.styles["Subtitle"]
    document.add_heading("Purpose", level=1)
    document.add_paragraph(
        "Use each scenario only as inspiration. Write a new, realistic complaint in all four "
        "language forms yourself. Do not copy the scenario brief as the English variant and do "
        "not use an LLM or machine translation to produce the variants."
    )
    document.add_heading("Rules", level=1)
    rules = (
        "Keep all four variants semantically equivalent: same failure, amount, date, and outcome.",
        "Make Roman Urdu naturally messy rather than standardized or transliterated word-for-word.",
        "Include at least one concrete fact in every variant.",
        "Copy that fact exactly into the Fact span field for the same variant.",
        "Use fictional values only. Do not enter real names, phone numbers, CNICs, or accounts.",
        "The 20 cases are a private development pilot, not publishable benchmark results.",
    )
    for rule in rules:
        document.add_paragraph(rule, style="List Bullet")
    document.add_paragraph(
        "Taxonomy note: the benchmark has 24 valid intents. This workbook proposes 20 distinct, "
        "high-value complaint cases. The appendix lists all 24 labels for reference."
    )


def _add_case_table(document: DocumentType) -> None:
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ("Language form", "Your human-written complaint", "Exact fact_span")
    for cell, text in zip(table.rows[0].cells, headers, strict=True):
        _set_cell_text(cell, text, bold=True)
        _shade_cell(cell, "DDEBE4")

    for language in LANGUAGE_FORMS:
        cells = table.add_row().cells
        _set_cell_text(cells[0], LANGUAGE_LABELS[language], bold=True)
        _set_cell_text(cells[1], "[Write your complaint here]\n\n\n")
        _set_cell_text(cells[2], "[Copy one exact fact here]\n\n")
        if language == "urdu_script":
            cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    widths = (Inches(1.35), Inches(4.75), Inches(1.35))
    for row in table.rows:
        for cell, width in zip(row.cells, widths, strict=True):
            cell.width = width


def _add_case(document: DocumentType, index: int, guide: CaseGuide, intent: Intent) -> None:
    if index > 1:
        document.add_section(WD_SECTION.NEW_PAGE)
    document.add_heading(f"Case {index:02d}: {guide.intent_id}", level=1)
    family = intent.family.replace("_", " ").title()
    metadata = document.add_paragraph()
    metadata.add_run("Family: ").bold = True
    metadata.add_run(family)
    metadata.add_run("    Gold intent: ").bold = True
    metadata.add_run(guide.intent_id)

    document.add_heading("Official definition", level=2)
    document.add_paragraph(intent.definition)
    document.add_heading("Complete example complaints", level=2)
    for example in intent.positive_examples:
        document.add_paragraph(example, style="List Bullet")
    document.add_heading("Inspiration scenario (do not copy)", level=2)
    document.add_paragraph(guide.scenario_brief)
    document.add_heading("Concrete fact to include", level=2)
    document.add_paragraph(guide.concrete_fact)
    document.add_heading("Boundary to preserve", level=2)
    document.add_paragraph(guide.boundary_note)
    document.add_heading("Your four variants", level=2)
    _add_case_table(document)
    document.add_paragraph(
        "Human check: same intent and facts in all four forms; natural phrasing; no real PII.",
    )


def _add_taxonomy_appendix(document: DocumentType, taxonomy: Taxonomy) -> None:
    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_heading("Appendix: All 24 Valid Intent Labels", level=1)
    document.add_paragraph(
        "The 20 proposed cases above use a focused subset. These are all valid labels in taxonomy "
        f"version {taxonomy.version}."
    )
    current_family = None
    for intent in taxonomy.intents:
        if intent.family != current_family:
            current_family = intent.family
            document.add_heading(current_family.replace("_", " ").title(), level=2)
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(f"{intent.id}: ").bold = True
        paragraph.add_run(intent.definition)


def render_workbook(taxonomy: Taxonomy, output: Path) -> None:
    """Create the formatted Word workbook at a private, ignored path."""
    intent_map = {intent.id: intent for intent in taxonomy.intents}
    missing = [guide.intent_id for guide in CASE_GUIDES if guide.intent_id not in intent_map]
    if missing:
        raise ValueError(f"Workbook case intents missing from taxonomy: {', '.join(missing)}")

    document = Document()
    _configure_document(document)
    _add_intro(document)
    for index, guide in enumerate(CASE_GUIDES, start=1):
        _add_case(document, index, guide, intent_map[guide.intent_id])
    _add_taxonomy_appendix(document, taxonomy)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def parse_args() -> argparse.Namespace:
    """Parse Word workbook generation options."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--taxonomy", type=Path, default=DEFAULT_TAXONOMY)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def main() -> None:
    """Load the taxonomy and write the private authoring workbook."""
    args = parse_args()
    render_workbook(load_taxonomy(args.taxonomy), args.output)
    print(f"Wrote {len(CASE_GUIDES)}-case authoring workbook to {args.output}")


if __name__ == "__main__":
    main()
